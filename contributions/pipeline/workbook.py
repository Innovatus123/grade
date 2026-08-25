"""Build the tracker workbook from scratch off the parquet. Idempotent - does not
depend on any prior workbook state."""
import pandas as pd, numpy as np, os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = os.path.expanduser("~/pafin/out")
XL = f"{OUT}/PA_Political_Contributions_Tracker.xlsx"
AS_OF = "2026-08-04"

d = pd.read_parquet(f"{OUT}/contributions_unified.parquet")
v = pd.read_parquet(f"{OUT}/donor_entities.parquet")
review = pd.read_csv(f"{OUT}/entity_review_queue.csv")

# Always re-attach entity IDs from the current resolver output - resolve.py
# re-issues IDs on every run, so any copy already on the transactions is stale.
d = d.drop(columns=[c for c in ("donor_entity_id", "entity_name") if c in d.columns])
d = d.join(v.set_index("donor_name_key")[["donor_entity_id", "entity_name"]],
           on="donor_name_key")
d["donor_entity_id"] = d.donor_entity_id.fillna("UNRESOLVED")
d["entity_name"] = d.entity_name.fillna(d.donor_name_raw)
d.to_parquet(f"{OUT}/contributions_unified.parquet", index=False)
for j, g in d.groupby("jurisdiction", observed=True):
    for y, gy in g.groupby("report_year", observed=True):
        gy.to_csv(f"{OUT}/contributions_{j}_{y}.csv", index=False)

live = d[d.is_primary & (d.donor_bucket != "EXCLUDE-ROLLUP")].copy()
for c in ("office", "party", "recipient_type"):
    live[c] = live[c].astype(str).replace({"nan": "", "None": ""}).fillna("")

# --------------------------------------------------------------------- README --
readme = pd.DataFrame({"Item": [], "Detail": []})
rows = [
 ("Purpose", "Track all contributions to Pennsylvania politicians, split organizational (corporate/PAC/union/party) vs individual."),
 ("Coverage window", "Filing years 2024, 2025, 2026 (partial - PA state export current through 2026-07-24)"),
 ("Data as of", AS_OF),
 ("Total tracked", f"${live.amount.sum():,.0f} across {len(live):,} contributions"),
 ("", ""),
 ("SOURCE 1", "PA Dept. of State Full Campaign Finance Export - annual ZIP: filer / contrib / expense / debt / receipt"),
 ("SOURCE 2", "City of Philadelphia campfin_contributions via Carto SQL API, per report_year"),
 ("SOURCE 3", "FEC bulk: cn (candidates), ccl (linkage), itpas2 (PAC-to-candidate), itcont (individuals), filtered to PA"),
 ("", ""),
 ("RULE - PA state", "Corporations MAY NOT contribute directly to candidates - 25 P.S. Sec. 3253. Corporate money must flow through a separate segregated fund (PAC). Exception: ballot questions."),
 ("RULE - PA state", "No dollar limits on contributions to PA state candidates beyond that treasury ban."),
 ("RULE - PA state", "Practical consequence: at state level corporate money is almost never labeled corporate. Tracing corporate influence means tracing PAC sponsors, not reading donor names."),
 ("RULE - Philadelphia", "Individual limit $3,700/yr; non-individual (PAC, partnership, sole proprietorship, other business organization) limit $14,800/yr. Effective 2024-01-01, next quadrennial CPI adjustment 2028."),
 ("RULE - Philadelphia", "Pay-to-play: exceeding the limit bars non-competitively bid city contracts (>$10,000 individual / >$25,000 business) for the officeholder's term. Attribution reaches affiliates, partners, directors, officers and associated PACs."),
 ("RULE - Federal", "Corporate treasury contributions to candidates prohibited; corporate PACs (SSFs) permitted. Limits indexed per two-year cycle - verify at fec.gov before compliance use."),
 ("", ""),
 ("CLASSIFICATION", "No jurisdiction publishes a reliable corporate/individual flag. PA state has no donor-type field at all. Philadelphia has one but it is 92% 'Not Specified'. Only the FEC entity code is reliable."),
 ("CLASSIFICATION", "All classification is derived. Every row carries donor_class, donor_class_method and donor_class_conf. Filter on confidence before quoting any figure externally."),
 ("DEDUPLICATION", "Philadelphia and PA state both carry contributions for committees that dual-file - 47% of Philadelphia dollars duplicated a PA-STATE row. Rows are flagged, not deleted. Take all totals on is_primary = TRUE."),
 ("ROLLUP EXCLUSION", "Report summary lines ('Unitemized', 'Anonymous', 'Federal Contributions Non-PA Activity', 'Non Pennsylvania Receipts') carried $83.1M and read as mega-donors until screened out. They are flagged donor_bucket = EXCLUDE-ROLLUP."),
 ("", ""),
 ("ENTITY RESOLUTION", f"Added {AS_OF}. {len(v):,} donor name variants resolved to {v.donor_entity_id.nunique():,} entities. Every transaction carries donor_entity_id."),
 ("ENTITY RESOLUTION", "No deterministic identifier exists in any source. Philadelphia publishes donor_id, donor_name_std and donor_filer_id but all three are 0% populated. Matching is derived from name structure plus zip and employer corroboration."),
 ("ENTITY RESOLUTION", "Persons: name parsed into first/middle/last/suffix, handling both 'LAST, FIRST' and 'FIRST LAST'. Conservative nickname expansion only - NANCY, MARIA, LINDA, TRACY, CAROL and similar are deliberately NOT treated as nicknames; doing so caused measurable over-merging."),
 ("ENTITY RESOLUTION", "Organizations: legal forms stripped and abbreviations expanded, but embedded numbers must match exactly. IBEW Local 98 and IBEW Local 5 can never merge."),
 ("ENTITY RESOLUTION", "An initial-only first name ('A WILLIAMS') never auto-merges alone - it requires two independent corroborations (matching zip AND matching employer). Otherwise it goes to review."),
 ("ENTITY RESOLUTION", "After clustering, each person cluster is re-partitioned by expanded first name, then again by middle initial and state, to break transitive chains. This split 3,224 clusters - it is what keeps a Reading PA John M. Arnold out of the Houston John D. Arnold entity."),
 ("ENTITY RESOLUTION", "Sheet 11 is the human review queue - pairs scoring 0.82-0.90 that were NOT merged. Work it top-down by dollar value."),
 ("", ""),
 ("LIMITATION 1", "Entities are deliberately under-merged rather than over-merged. A donor's resolved total is a floor, not a ceiling."),
 ("LIMITATION 2", "PA candidate FILERIDs are re-issued every election year (e.g. 2026C0421). A politician master crosswalk for cross-year tracking is NOT yet built."),
 ("LIMITATION 3", "Federal individual contributions cover the 2026 cycle file only. Earlier cycles need additional indiv[YY].zip pulls."),
 ("LIMITATION 4", "PA state export is filing-year based - a contribution can appear in a different report year than its transaction date."),
 ("", ""),
 ("REFRESH", "Run build.py, then resolve.py, then workbook.py. All three are idempotent. Runtime approximately 20 minutes, dominated by the 2GB FEC individual-contributions download."),
 ("REFRESH", "PA state: full-year files are restated periodically - re-pull {year}.zip. Philadelphia: near real-time. FEC: weekly, typically Sunday."),
 ("REFRESH", "Recommended cadence: monthly, plus within one week of each PA reporting deadline."),
]
readme = pd.DataFrame(rows, columns=["Item", "Detail"])

# ------------------------------------------------------------------- QA sheets --
qa = pd.DataFrame({
 "metric": ["Total rows ingested", "Primary (deduped) rows", "Cross-jurisdiction duplicates",
            "Rollup / reporting lines excluded", "Total $ (primary, ex-rollup)",
            "Individual $", "Organizational $", "Unresolved $",
            "Rows classified below 0.70 confidence", "$ classified below 0.70 confidence"],
 "value": [len(d), int(d.is_primary.sum()), int((~d.is_primary).sum()),
           int((d.donor_bucket == "EXCLUDE-ROLLUP").sum()), live.amount.sum(),
           live.loc[live.donor_bucket == "INDIVIDUAL", "amount"].sum(),
           live.loc[live.donor_bucket == "ORGANIZATIONAL", "amount"].sum(),
           live.loc[live.donor_bucket == "OTHER/UNRESOLVED", "amount"].sum(),
           int((live.donor_class_conf < 0.7).sum()),
           live.loc[live.donor_class_conf < 0.7, "amount"].sum()]})

gain = (v.groupby("donor_entity_id")
         .agg(variants=("raw", "size"), total=("total", "sum"),
              name=("entity_name", "first"), bucket=("bucket", "first"))
         .query("variants > 1").sort_values("total", ascending=False))
top100 = live.groupby("donor_entity_id", observed=True).amount.sum().nlargest(100).index
eqa = pd.DataFrame({
 "metric": ["Donor name variants (input)", "Resolved entities (output)", "Variants absorbed",
            "Reduction", "Deterministic merges", "Fuzzy merges (scored)",
            "Clusters split for first-name conflict", "Pairs queued for human review",
            "Multi-variant entities", "Largest entity (variants)",
            "Top-100 entities that were fragmented", "$ inside multi-variant entities"],
 "value": [len(v), v.donor_entity_id.nunique(), len(v) - v.donor_entity_id.nunique(),
           f"{1 - v.donor_entity_id.nunique()/len(v):.1%}", 30824, 8306, 3224, len(review),
           len(gain), int(v.groupby("donor_entity_id").size().max()),
           int(sum(1 for e in top100 if e in set(gain.index))), float(gain.total.sum())]})

# ------------------------------------------------------------------ analysis ----
def pivot(df, idx):
    p = df.pivot_table(index=idx, columns="donor_bucket", values="amount",
                       aggfunc="sum", fill_value=0, observed=True).reset_index()
    for c in ("INDIVIDUAL", "ORGANIZATIONAL", "OTHER/UNRESOLVED"):
        if c not in p: p[c] = 0.0
    p["TOTAL"] = p.INDIVIDUAL + p.ORGANIZATIONAL + p["OTHER/UNRESOLVED"]
    p["ORG_SHARE"] = (p.ORGANIZATIONAL / p.TOTAL.replace(0, np.nan)).round(3)
    return p.sort_values("TOTAL", ascending=False)

cand = live[live.recipient_type.str.contains("Candidate", na=False)]
cv = pivot(cand, ["jurisdiction", "office", "party", "recipient_name"])
allrec = pivot(live, ["jurisdiction", "office", "recipient_name", "party"])

ent = (live.groupby("donor_entity_id", observed=True)
       .agg(entity_name=("entity_name", "first"), donor_class=("donor_class", "first"),
            donor_bucket=("donor_bucket", "first"), total=("amount", "sum"),
            gifts=("amount", "size"), recipients=("recipient_name", "nunique"),
            jurisdictions=("jurisdiction", "nunique"), city=("donor_city", "first"),
            state=("donor_state", "first")).reset_index())
ent["variants"] = v.groupby("donor_entity_id").size().reindex(ent.donor_entity_id).fillna(1).values
ent = ent.sort_values("total", ascending=False)

sheets = {
 "00_README_and_Rules": readme,
 "01_QA_Control": qa,
 "02_Entity_Resolution_QA": eqa,
 "03_Candidates_Corp_vs_Indiv": cv,
 "04_All_Recipients": allrec,
 "05_Top_Org_Entities": ent[ent.donor_bucket == "ORGANIZATIONAL"].nlargest(600, "total"),
 "06_Top_Indiv_Entities": ent[ent.donor_bucket == "INDIVIDUAL"].nlargest(600, "total"),
 "07_Merged_Entities": gain.reset_index().head(800),
 "08_Multi_Recipient_Entities": ent[ent.recipients >= 5].nlargest(500, "total"),
 "09_Recipient_By_DonorClass": live.pivot_table(index=["jurisdiction", "recipient_name"],
        columns="donor_class", values="amount", aggfunc="sum", fill_value=0,
        observed=True).reset_index(),
 "10_Classification_Method": live.groupby("donor_class_method", observed=True)
        .agg(rows=("amount", "size"), dollars=("amount", "sum"))
        .sort_values("dollars", ascending=False).reset_index(),
 "11_Entity_Review_Queue": review.head(3000),
}
with pd.ExcelWriter(XL, engine="openpyxl") as w:
    for name, df in sheets.items():
        df.to_excel(w, sheet_name=name, index=False)

wb = load_workbook(XL)
HDR = PatternFill("solid", fgColor="1F3864")
for ws in wb.worksheets:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = HDR
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.iter_cols(min_row=1, max_row=1):
        h = str(col[0].value or "")
        ws.column_dimensions[get_column_letter(col[0].column)].width = \
            70 if h == "Detail" else min(max(len(h) + 4, 14), 42)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                h = str(ws.cell(1, cell.column).value or "")
                cell.number_format = ("0.0%" if "SHARE" in h.upper() else
                                      "0.000" if h == "score" else
                                      "#,##0" if abs(cell.value) >= 1000 else "#,##0.00")
wb.save(XL)
ent.to_csv(f"{OUT}/donor_entity_master.csv", index=False)
print(qa.to_string(index=False)); print()
print(eqa.to_string(index=False)); print()
print("sheets:", wb.sheetnames)
print("\nTop 12 org entities:")
print(sheets["05_Top_Org_Entities"].head(12)[["entity_name","total","gifts","recipients","variants"]].to_string(index=False))
print("\nTop 10 individual entities:")
print(sheets["06_Top_Indiv_Entities"].head(10)[["entity_name","total","gifts","recipients","variants"]].to_string(index=False))
